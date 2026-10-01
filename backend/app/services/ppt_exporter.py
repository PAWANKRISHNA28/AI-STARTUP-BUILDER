from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def generate_ppt_pitch_deck(blueprint_data: dict, output_path: str) -> str:
    prs = Presentation()
    
    # Set slide dimensions to widescreen 16:9
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    meta = blueprint_data.get("meta", {})
    title_text = meta.get("title", "Startup Pitch Deck")
    pitch = blueprint_data.get("investor_pitch", {})
    slides_data = pitch.get("slides", [])

    blank_slide_layout = prs.slide_layouts[6]

    for s_idx, s_data in enumerate(slides_data):
        slide = prs.slides.add_slide(blank_slide_layout)
        
        # Header / Title Text Box
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.733), Inches(1.0))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"Slide {s_data.get('slide')}: {s_data.get('title')}"
        p.font.size = Pt(32)
        p.font.bold = True
        p.font.color.rgb = RGBColor(79, 70, 229) # #4F46E5 Indigo

        # Content Text Box
        content_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.733), Inches(4.5))
        ctf = content_box.text_frame
        ctf.word_wrap = True
        cp = ctf.paragraphs[0]
        cp.text = s_data.get('content', '')
        cp.font.size = Pt(22)
        cp.font.color.rgb = RGBColor(55, 65, 81) # #374151 Dark Gray
        
        # Sub-bullet details if available
        if s_idx == 0:
            p2 = ctf.add_paragraph()
            p2.text = f"\nIndustry: {meta.get('industry')} | Target Users: {meta.get('target_users')}"
            p2.font.size = Pt(16)
            p2.font.color.rgb = RGBColor(107, 114, 128)

    prs.save(output_path)
    return output_path
