from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def generate_docx_report(blueprint_data: dict, output_path: str) -> str:
    doc = Document()
    
    meta = blueprint_data.get("meta", {})
    title = meta.get("title", "Startup Blueprint")
    idea = meta.get("idea", "")

    # Document Header
    p_title = doc.add_paragraph()
    run_title = p_title.add_run(f"AI STARTUP BUILDER - BLUEPRINT\n{title}")
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(79, 70, 229)
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT

    p_meta = doc.add_paragraph()
    r_meta = p_meta.add_run(f"Idea: {idea}\nIndustry: {meta.get('industry')} | Budget: {meta.get('budget')} | Target: {meta.get('target_users')}")
    r_meta.font.size = Pt(11)
    r_meta.font.italic = True
    r_meta.font.color.rgb = RGBColor(107, 114, 128)

    doc.add_heading('1. Executive Summary & Validation', level=1)
    idea_val = blueprint_data.get("idea_analyzer", {})
    doc.add_paragraph(f"Validation Score: {idea_val.get('validation_score')}/100")
    doc.add_paragraph(f"Problem Statement: {idea_val.get('problem_statement')}")
    doc.add_paragraph(f"Solution Overview: {idea_val.get('solution_overview')}")
    doc.add_paragraph(f"Unique Value Proposition: {idea_val.get('unique_value_proposition')}")

    doc.add_heading('2. Market Research (TAM / SAM / SOM)', level=1)
    mkt = blueprint_data.get("market_research", {})
    doc.add_paragraph(f"Total Addressable Market (TAM): {mkt.get('tam')}")
    doc.add_paragraph(f"Serviceable Addressable Market (SAM): {mkt.get('sam')}")
    doc.add_paragraph(f"Serviceable Obtainable Market (SOM): {mkt.get('som')}")
    doc.add_paragraph(f"Industry Growth Rate: {mkt.get('industry_growth_rate')}")

    doc.add_heading('3. Technical Architecture & Database Schema', level=1)
    db = blueprint_data.get("database_designer", {})
    doc.add_paragraph(f"Tables: {len(db.get('tables', []))} core entities generated.")
    doc.add_paragraph(f"SQL Schema:\n{db.get('sql_schema')}")

    doc.add_heading('4. Financial Estimates & ROI', level=1)
    fin = blueprint_data.get("finance", {})
    doc.add_paragraph(f"Development Cost: {fin.get('initial_build_cost')}")
    doc.add_paragraph(f"Monthly Hosting Infrastructure: {fin.get('monthly_server_infra')}")
    doc.add_paragraph(f"Break-Even Point: {fin.get('break_even_timeline')}")

    doc.save(output_path)
    return output_path
